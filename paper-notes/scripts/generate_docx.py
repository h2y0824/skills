"""
generate_docx.py — 生成论文阅读笔记 .docx 文件

用法:
    python generate_docx.py --input notes.json [--output filename.docx]

输入 JSON 格式:
{
    "title": "论文标题",
    "sections": [
        {
            "heading": "一、总体概述",
            "level": 1,
            "content": "正文内容..."
        },
        {
            "heading": "2.1 研究背景与动机",
            "level": 2,
            "content": "正文内容..."
        }
    ]
}
"""

import argparse
import json
import re
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_run_font(run, size_pt=12):
    """设置 run 的字体：中文宋体，英文 Times New Roman"""
    run.font.size = Pt(size_pt)
    run.font.name = "Times New Roman"
    # 设置中文字体
    r = run._element
    rPr = r.find(qn("w:rPr"))
    if rPr is None:
        rPr = r.makeelement(qn("w:rPr"), {})
        r.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), "宋体")


def set_paragraph_format(paragraph, indent_first_line=True, line_spacing=1.5):
    """设置段落格式：首行缩进、行距"""
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line_spacing
    if indent_first_line:
        fmt.first_line_indent = Cm(0.74)  # 约 2 个中文字符


def add_heading(doc, text, level=1):
    """添加标题"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Times New Roman"
        r = run._element
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            rPr = r.makeelement(qn("w:rPr"), {})
            r.insert(0, rPr)
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = rPr.makeelement(qn("w:rFonts"), {})
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:eastAsia"), "宋体")
    return heading


def add_paragraph(doc, text):
    """添加正文段落，支持简单的格式标记"""
    p = doc.add_paragraph()
    set_paragraph_format(p)

    # 处理简单的格式标记
    # **bold** -> bold
    # `code` -> code style
    parts = re.split(r"(\*\*.*?\*\*|`[^`]+`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run)
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run)
            run.font.name = "Consolas"
        else:
            run = p.add_run(part)
            set_run_font(run)

    return p


def add_bullet(doc, text):
    """添加列表项"""
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run)
    set_paragraph_format(p, indent_first_line=False)
    return p


def sanitize_filename(name, max_len=50):
    """清理文件名"""
    name = re.sub(r'[/\\:*?"<>|]', "", name)
    if len(name) > max_len:
        name = name[:max_len]
    return name


def main():
    parser = argparse.ArgumentParser(description="Generate paper notes .docx")
    parser.add_argument("--input", required=True, help="Path to notes JSON file")
    parser.add_argument("--output", default=None, help="Output .docx filename")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    title = data.get("title", "论文笔记")
    sections = data.get("sections", [])

    doc = Document()

    # 设置默认字体
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for section in sections:
        heading = section.get("heading", "")
        level = section.get("level", 1)
        content = section.get("content", "")

        add_heading(doc, heading, level=level)

        # 处理内容：按行分割，识别列表项和段落
        lines = content.strip().split("\n")
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if line.startswith("- ") or line.startswith("* "):
                add_bullet(doc, line[2:])
            elif line.startswith("• "):
                add_bullet(doc, line[2:])
            else:
                add_paragraph(doc, line)

    # 生成输出文件名
    if args.output:
        output_path = args.output
    else:
        safe_title = sanitize_filename(title)
        output_path = f"{safe_title}_笔记.docx"

    doc.save(output_path)
    print(f"笔记已保存至: {output_path}")


if __name__ == "__main__":
    main()
