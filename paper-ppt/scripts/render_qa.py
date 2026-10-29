# -*- coding: utf-8 -*-
"""把 PPTX 导出为 PNG + 4 宫格拼图，用于视觉质检。

用法:
    py render_qa.py deck.pptx --out _qa
    py render_qa.py deck.pptx --out _qa --slides 7,9,11
    py render_qa.py deck.pptx --out _qa --pdf
    py render_qa.py deck.pptx --out _qa --no-sheets

优先用本机 PowerPoint (COM)；失败时回退 LibreOffice + PyMuPDF。
每次只导出指定页；迭代时不要整册重导。
"""
import argparse
import glob
import os
import shutil
import subprocess
import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from PIL import Image, ImageDraw, ImageFont


def parse_slides(spec, n):
    if not spec:
        return list(range(1, n + 1))
    out = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            out.extend(range(int(a), min(int(b), n) + 1))
        elif part:
            out.append(int(part))
    return [i for i in out if 1 <= i <= n]


def slide_count(pptx):
    from pptx import Presentation
    return len(Presentation(pptx).slides)


def export_com(pptx, slides, outdir, width, height, pdf_path=None):
    import win32com.client
    app = win32com.client.Dispatch("PowerPoint.Application")
    deck = app.Presentations.Open(os.path.abspath(pptx),
                                  ReadOnly=True, WithWindow=False)
    try:
        for i in slides:
            deck.Slides(i).Export(
                os.path.abspath(os.path.join(outdir, "s%02d.png" % i)),
                "PNG", width, height)
        if pdf_path:
            deck.SaveAs(os.path.abspath(pdf_path), 32)  # 32 = ppSaveAsPDF
    finally:
        deck.Close()


def export_libreoffice(pptx, slides, outdir, width, height, pdf_path=None):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("neither PowerPoint COM nor LibreOffice is available")
    tmp = os.path.join(outdir, "_pdf")
    os.makedirs(tmp, exist_ok=True)
    subprocess.run([exe, "--headless", "--convert-to", "pdf",
                    "--outdir", tmp, os.path.abspath(pptx)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdfs = glob.glob(os.path.join(tmp, "*.pdf"))
    if not pdfs:
        raise RuntimeError("LibreOffice produced no PDF")
    pdf = pdfs[0]
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open(pdf)
    for i in slides:
        page = doc[i - 1]
        zoom = width / page.rect.width
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
        pix.save(os.path.join(outdir, "s%02d.png" % i))
    doc.close()
    if pdf_path:
        shutil.copyfile(pdf, pdf_path)


def load_label_font(size=18):
    for name in ("segoeui.ttf", "arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    return ImageFont.load_default()


def build_sheets(outdir, slides, tile_w=800, tile_h=450, cols=2):
    """4 宫格拼图；每格左上角标页码，方便对页。"""
    font = load_label_font(18)
    paths = []
    for start in range(0, len(slides), cols * 2):
        group = slides[start:start + cols * 2]
        rows = (len(group) + cols - 1) // cols
        pad = 12
        sheet = Image.new("RGB", (tile_w * cols + pad * (cols + 1),
                                  tile_h * rows + pad * (rows + 1)), "white")
        draw = ImageDraw.Draw(sheet)
        for k, num in enumerate(group):
            im = Image.open(os.path.join(outdir, "s%02d.png" % num))
            im = im.convert("RGB").resize((tile_w, tile_h), Image.LANCZOS)
            x = pad + (k % cols) * (tile_w + pad)
            y = pad + (k // cols) * (tile_h + pad)
            sheet.paste(im, (x, y))
            label = "s%02d" % num
            tb = draw.textbbox((0, 0), label, font=font)
            draw.rectangle([x + 6, y + 6, x + 6 + tb[2] + 10, y + 6 + tb[3] + 10],
                           fill=(255, 255, 255))
            draw.text((x + 11, y + 8), label, fill=(20, 40, 60), font=font)
        path = os.path.join(outdir, "sheet%d.png" % (start // (cols * 2) + 1))
        sheet.save(path)
        paths.append(path)
    return paths


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--out", default="_qa")
    ap.add_argument("--slides", default=None, help="如 1-24 或 7,9,11")
    ap.add_argument("--width", type=int, default=1600)
    ap.add_argument("--height", type=int, default=900)
    ap.add_argument("--pdf", action="store_true", help="同时导出 PDF")
    ap.add_argument("--no-sheets", action="store_true")
    ap.add_argument("--tile-width", type=int, default=800)
    args = ap.parse_args()

    os.makedirs(args.out, exist_ok=True)
    n = slide_count(args.pptx)
    slides = parse_slides(args.slides, n)
    pdf_path = os.path.splitext(os.path.abspath(args.pptx))[0] + ".pdf" if args.pdf else None

    try:
        export_com(args.pptx, slides, args.out, args.width, args.height, pdf_path)
        backend = "PowerPoint"
    except Exception as e:
        print("COM export failed (%s); falling back to LibreOffice" % e, file=sys.stderr)
        export_libreoffice(args.pptx, slides, args.out, args.width, args.height, pdf_path)
        backend = "LibreOffice"

    sheets = []
    if not args.no_sheets:
        sheets = build_sheets(args.out, slides, tile_w=args.tile_width,
                              tile_h=args.tile_width * 9 // 16)

    print("exported %d slides via %s -> %s/" % (len(slides), backend, args.out))
    for p in sheets:
        print("sheet:", p)
    if pdf_path:
        print("pdf:", pdf_path)


if __name__ == "__main__":
    main()
