# -*- coding: utf-8 -*-
"""pptkit — 论文汇报 PPT 的通用版式库（16:9, 13.333 x 7.5 in）。

用法：
    from pptkit import *
    new_deck("论文短名 · 会议 '26 · 组会汇报")
    s = new_slide("01 / 背景", "标题写结论，不写主题")
    bullets(s, 0.7, 1.8, 6, 3, ["...", "..."])
    takeaway_strip(s, "一句话结论")
    notes(s, "本页主张 / 过渡 / 一个预判提问")
    save_deck("xxx_组会汇报.pptx")

所有尺寸单位为英寸。改版式只改这里，不在 build_deck.py 里重复造轮子。
"""
import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
from pptx.oxml.ns import qn
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (
    XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE,
    XL_TICK_MARK,
)

# ------------------------------------------------------------------ palette
NAVY = RGBColor(0x0E, 0x22, 0x38)
INK = RGBColor(0x1B, 0x2A, 0x38)
BODY = RGBColor(0x3D, 0x4B, 0x59)
MUTED = RGBColor(0x7A, 0x88, 0x96)
TEAL = RGBColor(0x0E, 0x8F, 0x8F)
TEAL_DARK = RGBColor(0x0A, 0x6E, 0x6E)
TEAL_LIGHT = RGBColor(0xDD, 0xF0, 0xEF)
AMBER = RGBColor(0xC9, 0x7A, 0x1E)
AMBER_LIGHT = RGBColor(0xFB, 0xEC, 0xD8)
RED = RGBColor(0xB4, 0x43, 0x3F)
RED_LIGHT = RGBColor(0xF8, 0xE4, 0xE2)
BLUE = RGBColor(0x2E, 0x5E, 0x8C)
BLUE_LIGHT = RGBColor(0xE3, 0xEC, 0xF4)
LIGHT = RGBColor(0xF4, 0xF7, 0xF9)
BORDER = RGBColor(0xDD, 0xE5, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY_FILL = RGBColor(0xE9, 0xEF, 0xF4)
FAINT_LINE = RGBColor(0xE8, 0xEE, 0xF3)

# 语义：TEAL=方法/计算，AMBER=传输/代价，RED=风险/局限，BLUE=基线/对照
FONT_LAT = "Segoe UI"
FONT_CN = "Microsoft YaHei"
FONT_MONO = "Consolas"
FONT_MATH = "Cambria Math"

SW, SH = 13.333, 7.5
ML = 0.7
CW = 11.93

prs = None
PAGE = 0
FOOTER = ""


def new_deck(footer_text=""):
    """初始化 16:9 画布；footer_text 为空时页脚只显示页码。"""
    global prs, PAGE, FOOTER
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    PAGE = 0
    FOOTER = footer_text
    return prs


def set_footer(text):
    """后续 new_slide 的页脚文字；在 new_deck 之后调用。"""
    global FOOTER
    FOOTER = text


def save_deck(path, title=None, subject=None, author=None, comments=None):
    """保存并写文档属性。"""
    if prs is None:
        raise RuntimeError("new_deck() must be called first")
    cp = prs.core_properties
    if title is not None:
        cp.title = title
    if subject is not None:
        cp.subject = subject
    if author is not None:
        cp.author = author
    if comments is not None:
        cp.comments = comments
    prs.save(path)
    return path


# ------------------------------------------------------------------ text
def hexs(color):
    if isinstance(color, RGBColor):
        return str(color)
    return "%02X%02X%02X" % (color[0], color[1], color[2])


def style_run(run, size=13, bold=False, color=BODY, italic=False,
              latin=FONT_LAT, ea=FONT_CN, spc=None, baseline=None):
    """同时设置西文与中文字体，避免换机器跑版。"""
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = latin
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", ea), ("a:cs", ea)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)
    if spc is not None:
        rPr.set("spc", str(spc))
    if baseline is not None:
        rPr.set("baseline", str(baseline))


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tb, tf


def para(tf, first=False, align=PP_ALIGN.LEFT, before=0, after=0, line=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    p.line_spacing = line
    return p


def add_runs(p, segs, size=13, color=BODY, bold=False, latin=FONT_LAT):
    """segs: "text" 或 [("text", {"size":..,"bold":..,"color":..}), ...]"""
    if isinstance(segs, str):
        segs = [segs]
    for seg in segs:
        if isinstance(seg, str):
            text, kw = seg, {}
        else:
            text, kw = seg
        r = p.add_run()
        r.text = text
        style_run(
            r,
            size=kw.get("size", size),
            bold=kw.get("bold", bold),
            color=kw.get("color", color),
            italic=kw.get("italic", False),
            latin=kw.get("latin", latin),
            spc=kw.get("spc"),
            baseline=kw.get("baseline"),
        )
    return p


def set_bullet(p, color=TEAL, char="•"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "228600")
    pPr.set("indent", "-228600")
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont", "a:buSzPct"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = buClr.makeelement(qn("a:srgbClr"), {"val": hexs(color)})
    buClr.append(srgb)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(buClr)
    pPr.append(buFont)
    pPr.append(buChar)


def bullets(slide, x, y, w, h, items, size=13, color=BODY, gap=8,
            bullet_color=TEAL, line=1.12, anchor=MSO_ANCHOR.TOP):
    tb, tf = textbox(slide, x, y, w, h, anchor=anchor)
    for i, item in enumerate(items):
        p = para(tf, first=(i == 0), before=(0 if i == 0 else gap), line=line)
        set_bullet(p, color=bullet_color)
        add_runs(p, item, size=size, color=color)
    return tb


def plain_paras(slide, x, y, w, h, lines, size=13, color=BODY, gap=6, line=1.15,
                anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    tb, tf = textbox(slide, x, y, w, h, anchor=anchor)
    for i, ln in enumerate(lines):
        p = para(tf, first=(i == 0), before=(0 if i == 0 else gap), line=line, align=align)
        add_runs(p, ln, size=size, color=color)
    return tb


# ------------------------------------------------------------------ shapes
def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0,
         shape=MSO_SHAPE.RECTANGLE, radius=None, dash=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
        if dash is not None:
            sp.line.dash_style = dash
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        sp.adjustments[0] = radius
    return sp


def card(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=0.06):
    return rect(slide, x, y, w, h, fill=fill, line=line, line_w=0.9,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)


def box(slide, x, y, w, h, text, fill=WHITE, line=BORDER, color=INK, size=11,
        bold=False, radius=0.09, align=PP_ALIGN.CENTER, dash=None, margins=0.05):
    sp = rect(slide, x, y, w, h, fill=fill, line=line, line_w=0.9,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius, dash=dash)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(margins)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 0.95
        r = p.add_run()
        r.text = ln
        style_run(r, size=size, bold=bold, color=color)
    return sp


def pill(slide, x, y, w, h, text, fill=TEAL_LIGHT, color=TEAL_DARK, size=9,
         bold=True, line=None):
    sp = rect(slide, x, y, w, h, fill=fill, line=line,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    style_run(r, size=size, bold=bold, color=color)
    return sp


def line(slide, x1, y1, x2, y2, color=BORDER, width=1.0, dash=None, head=False,
         tail=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash is not None:
        c.line.dash_style = dash
    ln = c.line._get_or_add_ln()
    if head:
        ln.append(ln.makeelement(qn("a:headEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))
    if tail:
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))
    return c


# ------------------------------------------------------------------ slide scaffold
def header(slide, kicker, title, title_size=24):
    rect(slide, ML, 0.46, 0.09, 0.09, fill=TEAL)
    tb, tf = textbox(slide, 0.9, 0.38, 11.4, 0.26)
    p = para(tf, first=True)
    r = p.add_run()
    r.text = kicker
    style_run(r, size=9.5, bold=True, color=TEAL, spc=160)
    tb, tf = textbox(slide, ML, 0.66, 12.0, 0.7)
    p = para(tf, first=True, line=0.95)
    r = p.add_run()
    r.text = title
    style_run(r, size=title_size, bold=True, color=INK)
    line(slide, ML, 1.47, ML + CW, 1.47, color=FAINT_LINE, width=0.9)


def footer(slide, pageno):
    line(slide, ML, 7.08, ML + CW, 7.08, color=FAINT_LINE, width=0.6)
    if FOOTER:
        tb, tf = textbox(slide, ML, 7.12, 9.8, 0.2)
        p = para(tf, first=True)
        r = p.add_run()
        r.text = FOOTER
        style_run(r, size=8.5, color=MUTED)
    tb, tf = textbox(slide, 11.6, 7.12, 1.03, 0.2)
    p = para(tf, first=True, align=PP_ALIGN.RIGHT)
    r = p.add_run()
    r.text = str(pageno)
    style_run(r, size=8.5, color=MUTED)


def new_slide(kicker=None, title=None, page=True, title_size=24):
    global PAGE
    if prs is None:
        raise RuntimeError("call new_deck() first")
    PAGE += 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    if kicker or title:
        header(s, kicker or "", title or "", title_size=title_size)
    if page:
        footer(s, PAGE)
    return s


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ------------------------------------------------------------------ content blocks
def takeaway_strip(slide, text, y=6.02, x=ML, w=CW, h=0.62,
                   label="一句话", fill=TEAL_LIGHT, bar=TEAL,
                   label_color=TEAL_DARK, text_color=BODY, size=10):
    """底部结论条：全页唯一结论，别写第二件事。"""
    rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    rect(slide, x, y, 0.06, h, fill=bar)
    tb, tf = textbox(slide, x + 0.25, y + 0.1, w - 0.38, h - 0.18)
    p = para(tf, first=True, line=1.12)
    segs = []
    if label:
        segs.append((label + "：", {"bold": True, "color": label_color, "size": size}))
    if isinstance(text, str):
        segs.append((text, {"color": text_color, "size": size}))
    else:
        segs.extend(text)
    add_runs(p, segs, size=size)


def metric_card(slide, x, y, w, h, value, label, sub=None, color=TEAL,
                fill=WHITE, line_color=BORDER, value_size=32):
    """大数字卡：value 用大号色字，label 解释口径，sub 放补充。"""
    card(slide, x, y, w, h, fill=fill, line=line_color)
    tb, tf = textbox(slide, x + 0.18, y + 0.12, w - 0.36, h - 0.2)
    p = para(tf, first=True)
    add_runs(p, [(value, {"size": value_size, "bold": True, "color": color})])
    p = para(tf, before=2, line=1.0)
    add_runs(p, [(label, {"size": 9.5, "color": MUTED})])
    if sub:
        p = para(tf, before=2, line=1.05)
        add_runs(p, [(sub, {"size": 8.5, "color": MUTED, "italic": True})])


def numbered_steps(slide, x, y, w, steps, step_w=None, gap_color=MUTED,
                   box_h=0.55, size=7.8):
    """横向流程条：steps 为字符串列表，自动画箭头；默认按总宽均分。"""
    steps = list(steps)
    if step_w is None:
        step_w = w / max(len(steps), 1)
    for i, st in enumerate(steps):
        bx = x + i * step_w
        last = (i == len(steps) - 1)
        box(slide, bx, y, step_w - 0.17, box_h, st,
            fill=AMBER_LIGHT if last else LIGHT,
            line=AMBER if last else BORDER,
            color=AMBER if last else BODY, size=size)
        if not last:
            line(slide, bx + step_w - 0.15, y + box_h / 2,
                 bx + step_w - 0.02, y + box_h / 2,
                 color=gap_color, width=0.8, tail=True)


# ------------------------------------------------------------------ table / chart
def styled_table(slide, x, y, w, h, rows, col_ratios=None, font_size=10,
                 header_fill=NAVY, header_color=WHITE, align=PP_ALIGN.CENTER,
                 first_col_align=None, row_heights=None):
    nrows, ncols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = gf.table
    tblPr = table._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for e in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(e)
    sid = tblPr.makeelement(qn("a:tableStyleId"), {})
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    tblPr.append(sid)
    if col_ratios is None:
        col_ratios = [1] * ncols
    tot = sum(col_ratios)
    for i, cr in enumerate(col_ratios):
        table.columns[i].width = Inches(w * cr / tot)
    if row_heights:
        for i, rh in enumerate(row_heights):
            table.rows[i].height = Inches(rh)
    else:
        for r in range(nrows):
            table.rows[r].height = Inches(h / nrows)
    for r in range(nrows):
        for c in range(ncols):
            spec = rows[r][c]
            if isinstance(spec, dict):
                text = spec.get("t", "")
                fill = spec.get("fill", WHITE if r % 2 else LIGHT)
                color = spec.get("color", BODY)
                bold = spec.get("bold", False)
                size = spec.get("size", font_size)
                al = spec.get("align", align)
            else:
                text = str(spec)
                fill = header_fill if r == 0 else (WHITE if r % 2 else LIGHT)
                color = header_color if r == 0 else BODY
                bold = (r == 0)
                size = font_size
                al = align
                if r > 0 and c == 0 and first_col_align is not None:
                    al = first_col_align
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = al
            p.line_spacing = 0.95
            r_ = p.add_run()
            r_.text = text
            style_run(r_, size=size, bold=bold, color=color)
    return table


def chart_font(chart, size=10, color=BODY):
    chart.font.name = FONT_LAT
    chart.font.size = Pt(size)
    chart.font.color.rgb = color


def style_axes(chart, ymin=None, ymax=None, major=None, fmt="0.0", grid=True):
    va = chart.value_axis
    va.has_major_gridlines = grid
    if grid:
        va.major_gridlines.format.line.color.rgb = RGBColor(0xEC, 0xF1, 0xF5)
        va.major_gridlines.format.line.width = Pt(0.6)
    va.format.line.color.rgb = BORDER
    va.format.line.width = Pt(0.6)
    va.tick_labels.font.size = Pt(9)
    va.tick_labels.font.color.rgb = MUTED
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    if ymin is not None:
        va.minimum_scale = ymin
    if ymax is not None:
        va.maximum_scale = ymax
    if major is not None:
        va.major_unit = major
    ca = chart.category_axis
    ca.format.line.color.rgb = BORDER
    ca.tick_labels.font.size = Pt(9.5)
    ca.tick_labels.font.color.rgb = BODY
    ca.major_tick_mark = XL_TICK_MARK.NONE      # 传 0 会报 ValueError
    ca.format.line.width = Pt(0.6)
